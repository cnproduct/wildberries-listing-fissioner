"""Offline regression checks. No WB/Feishu network requests."""
import copy
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import fission_exporter as exporter
import feishu_bot_service as bot


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        Image.new("RGB", (700, 900), "white").save(self.root / "source.jpg")
        self.data = {"brand": "Example", "category": "Термокружки", "composition": "Сталь",
                     "country": "Китай", "length_cm": 22, "width_cm": 9, "height_cm": 9,
                     "weight_kg": 0.35, "items": [
                         {"persona": f"Scenario {i}", "title": f"Термокружка вариант {i}",
                          "description": "Термокружка для напитков.", "keywords": ["термокружка", "чай"],
                          "barcode": "1234567890", "image_path": "source.jpg"} for i in range(1, 11)]}

    def export(self, **kwargs):
        return Path(exporter.build_wb_batch_excel(self.data, self.root / "out", "TEST", base_dir=self.root, **kwargs))

    def test_strict_round_trip_portable_images_and_literal_text(self):
        self.data["items"][0]["description"] = '=HYPERLINK("https://example.com")'
        path = self.export(strict=True)
        book = load_workbook(path)
        self.assertEqual(book["Listing_Review"].max_row, 11)
        cell = book["Listing_Review"]["D2"]
        self.assertEqual(cell.value, self.data["items"][0]["description"])
        self.assertEqual(cell.data_type, "s")
        manifest = json.loads((path.parent / "wb_fission_TEST_manifest.json").read_text())
        self.assertEqual(manifest["issues"], [])
        self.assertEqual(manifest["status"], "validated_assets")
        self.assertFalse(manifest["published"])
        for item in manifest["items"]:
            self.assertFalse(Path(item["image_path"]).is_absolute())
            with Image.open(path.parent / item["image_path"]) as img:
                self.assertEqual(img.format, "JPEG")
        self.assertIsNone(manifest["items"][0]["price"])
        self.assertIsNone(manifest["items"][0]["discount"])

    def test_each_missing_package_fact_is_not_fabricated(self):
        for field in exporter.PACKAGE_FIELDS:
            with self.subTest(field=field):
                payload = copy.deepcopy(self.data)
                del payload[field]
                rows, issues = exporter.validate_payload(payload, "TEST", self.root)
                self.assertIsNone(rows[0][field])
                self.assertTrue(any(i["field"] == field for i in issues))
                with self.assertRaises(ValueError):
                    exporter.build_wb_batch_excel(payload, self.root / field, "TEST", base_dir=self.root, strict=True)
                self.assertFalse((self.root / field).exists())

    def test_invalid_measurements_rejected_even_in_drafts(self):
        for field in exporter.PACKAGE_FIELDS:
            for value in (0, -1, float("nan"), float("inf"), True, "20"):
                with self.subTest(field=field, value=value):
                    payload = copy.deepcopy(self.data)
                    payload[field] = value
                    with self.assertRaises(ValueError):
                        exporter.validate_payload(payload, "TEST", self.root)

    def test_explicit_invalid_override_does_not_fall_back(self):
        self.data["items"][0]["weight_kg"] = 0
        with self.assertRaises(ValueError):
            self.export()
        self.assertFalse((self.root / "out").exists())

    def test_draft_missing_images_and_barcode_remain_missing(self):
        for item in self.data["items"]:
            del item["image_path"]
            del item["barcode"]
        path = self.export()
        manifest = json.loads((path.parent / "wb_fission_TEST_manifest.json").read_text())
        self.assertEqual(manifest["status"], "draft")
        self.assertEqual(len(manifest["issues"]), 20)
        self.assertIsNone(manifest["items"][0]["barcode"])
        self.assertIsNone(manifest["items"][0]["image_path"])
        self.assertEqual(list((path.parent / "images").iterdir()), [])

    def test_remote_image_is_unverified(self):
        for item in self.data["items"]:
            item.pop("image_path")
            item["image_url"] = "https://example.com/main.jpg"
        with self.assertRaises(ValueError):
            self.export(strict=True)

    def test_teaching_data_never_becomes_validated_assets(self):
        self.data["example_only"] = True
        with self.assertRaises(ValueError):
            self.export(strict=True)
        path = self.export()
        manifest = json.loads((path.parent / "wb_fission_TEST_manifest.json").read_text())
        self.assertTrue(manifest["example_only"])
        self.assertEqual(manifest["status"], "draft")

    def test_missing_corrupt_and_small_images_block_strict(self):
        (self.root / "broken.png").write_bytes(b"not an image")
        Image.new("RGB", (100, 100)).save(self.root / "small.png")
        for name in ("missing.png", "broken.png", "small.png"):
            with self.subTest(name=name):
                self.data["items"][0]["image_path"] = name
                with self.assertRaises(ValueError):
                    self.export(strict=True)
        self.assertFalse((self.root / "out").exists())

    def test_product_facts_cannot_conflict(self):
        self.data["items"][0]["composition"] = "Plastic"
        with self.assertRaises(ValueError):
            self.export(strict=True)
        # Same invariant when facts only appear on items.
        self.data.pop("composition")
        for item in self.data["items"]:
            item["composition"] = "Steel"
        self.data["items"][0]["composition"] = "Plastic"
        with self.assertRaises(ValueError):
            self.export(strict=True)

    def test_titles_count_and_sku_paths(self):
        for sku in ("../escape", "/tmp/escape", "", "a" * 65):
            with self.subTest(sku=sku), self.assertRaises(ValueError):
                exporter.validate_payload(self.data, sku, self.root)
        self.data["items"].pop()
        with self.assertRaises(ValueError):
            self.export()
        self.data["items"] = self.data["items"][:1]
        self.data["items"][0]["title"] = "я" * 61
        with self.assertRaises(ValueError):
            self.export(strict=True, expected_count=1)

    def test_no_overwrite_or_partial_package(self):
        path = self.export(strict=True)
        before = path.read_bytes()
        with self.assertRaises(FileExistsError):
            self.export(strict=True)
        self.assertEqual(path.read_bytes(), before)
        shutil_error = OSError("copy failed")
        with patch.object(exporter.shutil, "copy2", side_effect=shutil_error):
            with self.assertRaises(OSError):
                exporter.build_wb_batch_excel(self.data, self.root / "failed", "TEST", base_dir=self.root)
        self.assertEqual(list((self.root / "failed").iterdir()), [])

    def test_cli_resolves_images_relative_to_json(self):
        (self.root / "data.json").write_text(json.dumps(self.data), encoding="utf-8")
        result = subprocess.run([sys.executable, str(ROOT / "scripts/fission_exporter.py"), "--json", str(self.root / "data.json"), "--sku", "CLI",
                                 "--out", str(self.root / "cli"), "--strict"], cwd=ROOT, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue((self.root / "cli/wb_fission_CLI/images/CLI_V01_main.jpg").exists())


class LocalBotTests(unittest.TestCase):
    def test_identifiers_are_distinct(self):
        self.assertEqual(bot.resolve_product("211832049"), "211832049")
        self.assertEqual(bot.resolve_product("货号: 211832049"), "211832049")
        for text in ("SKU: 211832049", "barcode: 211832049", "https://wildberries.ru.evil.test/catalog/211832049/detail.aspx"):
            self.assertIsNone(bot.resolve_product(text))

    def test_receipt_does_not_promise_generation_or_save_tokens(self):
        self.assertIn("尚未接通", bot.build_local_reply("211832049"))
        self.assertIn("没有保存令牌", bot.build_local_reply("绑定API: dummy"))
        self.assertIn("没有持久化执行任务", bot.build_local_reply("进度"))


if __name__ == "__main__":
    unittest.main()
